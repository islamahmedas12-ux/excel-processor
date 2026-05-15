"""
Dynamic API runner (v2 configs).

Pipeline:  resolve bindings  ->  write cells (multi-sheet, byte-preserving)
           ->  LibreOffice recalc  ->  read output cells.

No formula is computed here; the calc-service (LibreOffice) does that.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from . import binding_resolver
from .recalc_service import recalc_xlsx
from .xlsx_patch import write_cells


def run(
    file_content: bytes,
    filename: str,
    config: dict,
    params: dict | None = None,
    manual_inputs: dict | None = None,
) -> dict[str, Any]:
    """
    Execute a v2-configured file and return its computed outputs.

    Returns { "success": True, "outputs": { "<cell or sheet!cell>": value } }.
    """
    writes = binding_resolver.resolve(config, params, manual_inputs)

    # Apply writes one sheet at a time; each call preserves all other bytes
    # (images, charts, other sheets) so chaining is safe.
    patched = file_content
    for sheet_name, cell_values in writes.items():
        if cell_values:
            patched = write_cells(patched, cell_values, sheet_name)

    recalculated = recalc_xlsx(patched, filename)

    wb = load_workbook(BytesIO(recalculated), data_only=True, keep_vba=False)
    try:
        outputs: dict[str, Any] = {}
        for out in config.get("outputs", []):
            cell = out["cell"]
            sheet = out.get("sheet")
            ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
            # Key by "Sheet!Cell" when a sheet is specified, else just "Cell",
            # so callers with multi-sheet outputs get unambiguous keys.
            key = f"{sheet}!{cell}" if sheet else cell
            outputs[key] = ws[cell].value
    finally:
        wb.close()

    return {"success": True, "outputs": outputs}
