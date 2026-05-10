"""
Image & QR code injection into Excel files using openpyxl.

insert_image  — place any PNG/JPG bytes into a worksheet at a given cell anchor
insert_qr     — generate a QR code and insert it (optionally embedding the verify URL)
insert_barcode — generate a barcode and embed it into an Excel file at a given cell
"""
from __future__ import annotations

import io
from typing import Optional


# ── QR code generation ───────────────────────────────────────────────────────

def _make_qr_png(data: str, box_size: int = 6, border: int = 2) -> bytes:
    import qrcode
    from qrcode.image.pil import PilImage

    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=box_size,
        border=border,
    )
    qr.add_data(data)
    qr.make(fit=True)
    img = qr.make_image(image_factory=PilImage)
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    return buf.getvalue()


# ── Barcode generation ───────────────────────────────────────────────────────

def _make_barcode_png(data: str) -> bytes:
    import barcode
    from barcode.writer import ImageWriter

    code = barcode.get_barcode_class('code128')
    bc = code(data, writer=ImageWriter())
    buf = io.BytesIO()
    bc.write(buf)
    return buf.getvalue()


# ── Excel injection (openpyxl) ───────────────────────────────────────────────

def _load_wb(excel_bytes: bytes):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(excel_bytes))


def _save_wb(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _get_ws(wb, sheet_name: Optional[str] = None):
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}")
        return wb[sheet_name]
    return wb.active


def insert_image(
    excel_bytes: bytes,
    image_bytes: bytes,
    cell: str = 'A1',
    sheet_name: Optional[str] = None,
    width_px: int = 120,
    height_px: int = 120,
) -> bytes:
    """
    Embed image_bytes (PNG/JPG) into an Excel file at the given cell.
    Returns the modified Excel bytes.
    """
    from openpyxl.drawing.image import Image as XlImage

    wb = _load_wb(excel_bytes)
    ws = _get_ws(wb, sheet_name)

    img = XlImage(io.BytesIO(image_bytes))
    img.width  = width_px
    img.height = height_px
    img.anchor = cell

    ws.add_image(img)
    return _save_wb(wb)


def insert_qr(
    excel_bytes: bytes,
    qr_data: str,
    cell: str = 'A1',
    sheet_name: Optional[str] = None,
    size_px: int = 120,
) -> bytes:
    """
    Generate a QR code for qr_data and embed it into the Excel file at cell.
    Returns the modified Excel bytes.
    """
    qr_png = _make_qr_png(qr_data)
    return insert_image(excel_bytes, qr_png, cell, sheet_name, size_px, size_px)


def insert_barcode(
    excel_bytes: bytes,
    barcode_data: str,
    cell: str = 'A1',
    sheet_name: Optional[str] = None,
) -> bytes:
    """
    Generate a barcode for barcode_data and embed it into the Excel file at cell.
    Returns the modified Excel bytes.
    """
    barcode_png = _make_barcode_png(barcode_data)
    return insert_image(excel_bytes, barcode_png, cell, sheet_name, 200, 60)
