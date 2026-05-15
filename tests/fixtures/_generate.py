"""
tests/fixtures/_generate.py
Generates the .xlsx fixtures for test_recalc.py using openpyxl.

These fixtures deliberately use REAL Excel functions (IF, VLOOKUP, AVERAGE,
ROUND, MAX/MIN, nested AND, multi-sheet refs) — the exact things a hand-rolled
Python evaluator cannot do. openpyxl writes formulas as strings with NO cached
values, so a passing test proves the LibreOffice calc-service actually computed
them.

Run: python tests/fixtures/_generate.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

FIXTURE_DIR = Path(__file__).parent


def make_formula_test():
    """
    Sheet1 inputs A1..A4; rich formulas a toy evaluator could never handle.
      B1 = SUM(A1:A4)
      B2 = AVERAGE(A1:A4)
      B3 = IF(B1>50,"high","low")
      B4 = ROUND(B1/3,2)
      B5 = MAX(A1:A4)-MIN(A1:A4)
      C1 = IF(AND(A1>5,A2>5),A1*A2,0)      (nested AND inside IF)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 10
    ws["A2"] = 20
    ws["A3"] = 30
    ws["A4"] = 40
    ws["B1"] = "=SUM(A1:A4)"
    ws["B2"] = "=AVERAGE(A1:A4)"
    ws["B3"] = '=IF(B1>50,"high","low")'
    ws["B4"] = "=ROUND(B1/3,2)"
    ws["B5"] = "=MAX(A1:A4)-MIN(A1:A4)"
    ws["C1"] = "=IF(AND(A1>5,A2>5),A1*A2,0)"
    wb.save(FIXTURE_DIR / "formula_test.xlsx")
    print("Created formula_test.xlsx")


def make_vlookup():
    """
    Rates sheet is a lookup table; Calc!B1 does a VLOOKUP into it.
      Calc!A1 = lookup key (input)
      Calc!B1 = VLOOKUP(A1, Rates!A2:B4, 2, FALSE)
      Calc!C1 = B1 * 10
    """
    wb = Workbook()
    rates = wb.active
    rates.title = "Rates"
    rates["A1"] = "Item"
    rates["B1"] = "Price"
    rates["A2"] = "apple"
    rates["B2"] = 3
    rates["A3"] = "banana"
    rates["B3"] = 2
    rates["A4"] = "cherry"
    rates["B4"] = 5

    calc = wb.create_sheet("Calc")
    calc["A1"] = "banana"
    calc["B1"] = "=VLOOKUP(A1,Rates!A2:B4,2,FALSE)"
    calc["C1"] = "=B1*10"

    wb.save(FIXTURE_DIR / "formula_vlookup.xlsx")
    print("Created formula_vlookup.xlsx")


def make_cross_sheet():
    """
    Same cell coordinate (A1) on two different sheets, summed on a third.
    A naive 'strip the sheet name' approach collides and returns the wrong
    value — this fixture only passes with a real engine.
      Data!A1  = 5
      Other!A1 = 99
      Calc!B1  = Data!A1 + Other!A1   -> 104
      Calc!C1  = (Data!A1 + Other!A1) * 4  -> 416
    """
    wb = Workbook()
    data = wb.active
    data.title = "Data"
    data["A1"] = 5

    other = wb.create_sheet("Other")
    other["A1"] = 99

    calc = wb.create_sheet("Calc")
    calc["B1"] = "=Data!A1+Other!A1"
    calc["C1"] = "=(Data!A1+Other!A1)*4"

    wb.save(FIXTURE_DIR / "formula_cross_sheet.xlsx")
    print("Created formula_cross_sheet.xlsx")


def make_formula_with_chart():
    """A workbook with a chart plus a formula cell (chart must survive recalc)."""
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A2"] = 2
    ws["A3"] = 3
    ws["B1"] = "=SUM(A1:A3)"

    ref = Reference(ws, min_col=1, min_row=1, max_col=1, max_row=3)
    chart = BarChart()
    chart.add_data(ref)
    ws.add_chart(chart, "D1")

    wb.save(FIXTURE_DIR / "formula_with_chart.xlsx")
    print("Created formula_with_chart.xlsx")


if __name__ == "__main__":
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    make_formula_test()
    make_vlookup()
    make_cross_sheet()
    make_formula_with_chart()
    print("All fixtures generated.")
