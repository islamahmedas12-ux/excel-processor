"""
tests/test_recalc.py
Integration tests for the recalc pipeline.

Formula calculation is delegated to the calc-service (LibreOffice). These
tests therefore require calc-service to be reachable (it is, inside
docker-compose). If it is unreachable the tests skip with a clear message
rather than producing a misleading green/red.

The fixtures use REAL Excel functions (IF, VLOOKUP, AVERAGE, ROUND, MAX/MIN,
nested AND, multi-sheet refs). openpyxl writes them with no cached value, so a
pass proves LibreOffice actually computed them — not a hand-rolled evaluator.
"""

import io
import os
from pathlib import Path

import pytest
import requests
from openpyxl import load_workbook

from backend.services.excel_service import ExcelService
from backend.services.recalc_service import recalc_xlsx, CALC_SERVICE_URL

FIXTURE_DIR = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="session", autouse=True)
def _require_calc_service():
    """Skip the whole module if calc-service is not reachable."""
    try:
        r = requests.get(f"{CALC_SERVICE_URL}/health", timeout=5)
        if r.status_code != 200:
            pytest.skip(f"calc-service unhealthy at {CALC_SERVICE_URL}")
    except requests.RequestException:
        pytest.skip(
            f"calc-service unreachable at {CALC_SERVICE_URL} "
            f"(run via docker compose to exercise these tests)"
        )


@pytest.fixture
def formula_file():
    return (FIXTURE_DIR / "formula_test.xlsx").read_bytes()


@pytest.fixture
def vlookup_file():
    return (FIXTURE_DIR / "formula_vlookup.xlsx").read_bytes()


@pytest.fixture
def cross_sheet_file():
    return (FIXTURE_DIR / "formula_cross_sheet.xlsx").read_bytes()


@pytest.fixture
def chart_file():
    return (FIXTURE_DIR / "formula_with_chart.xlsx").read_bytes()


# ── recalc_xlsx (LibreOffice via calc-service) ────────────────────────────────

class TestRecalcService:
    def test_returns_bytes(self, formula_file):
        out = recalc_xlsx(formula_file, "formula_test.xlsx")
        assert isinstance(out, bytes) and len(out) > 0

    def test_formula_cells_get_cached_values(self, formula_file):
        out = recalc_xlsx(formula_file, "formula_test.xlsx")
        ws = load_workbook(io.BytesIO(out), data_only=True).active
        # A1..A4 = 10,20,30,40
        assert ws["B1"].value == 100          # SUM
        assert ws["B2"].value == 25           # AVERAGE
        assert ws["B3"].value == "high"       # IF (string result)
        assert round(ws["B4"].value, 2) == 33.33   # ROUND(100/3,2)
        assert ws["B5"].value == 30           # MAX-MIN
        assert ws["C1"].value == 200          # IF(AND(...),A1*A2,0)

    def test_literals_preserved(self, formula_file):
        out = recalc_xlsx(formula_file, "formula_test.xlsx")
        ws = load_workbook(io.BytesIO(out), data_only=True).active
        assert ws["A1"].value == 10
        assert ws["A4"].value == 40


# ── execute() end-to-end (write inputs → recalc → read outputs) ───────────────

class TestExecuteRealFormulas:
    def test_rich_formula_sheet(self, formula_file):
        res = ExcelService().execute(
            file_content=formula_file,
            filename="formula_test.xlsx",
            inputs={"A1": 5, "A2": 5, "A3": 5, "A4": 5},
            outputs=["B1", "B2", "B3", "B4", "B5", "C1"],
        )
        assert res["success"] is True
        r = res["results"]
        assert r["B1"] == 20                 # SUM(5,5,5,5)
        assert r["B2"] == 5                  # AVERAGE
        assert r["B3"] == "low"              # IF(20>50,...) -> low
        assert round(r["B4"], 2) == 6.67     # ROUND(20/3,2)
        assert r["B5"] == 0                  # MAX-MIN (all equal)
        # A1>5 is False when A1=5, so AND(...) is False -> IF returns 0
        assert r["C1"] == 0

    def test_vlookup(self, vlookup_file):
        res = ExcelService().execute(
            file_content=vlookup_file,
            filename="formula_vlookup.xlsx",
            inputs={"A1": "cherry"},
            outputs=["B1", "C1"],
            sheet_name="Calc",
        )
        assert res["success"] is True
        assert res["results"]["B1"] == 5     # VLOOKUP cherry -> 5
        assert res["results"]["C1"] == 50    # 5 * 10

    def test_cross_sheet_no_collision(self, cross_sheet_file):
        res = ExcelService().execute(
            file_content=cross_sheet_file,
            filename="formula_cross_sheet.xlsx",
            inputs={},
            outputs=["B1", "C1"],
            sheet_name="Calc",
        )
        assert res["success"] is True
        assert res["results"]["B1"] == 104   # Data!A1(5) + Other!A1(99)
        assert res["results"]["C1"] == 416   # *4

    def test_chart_preserved(self, chart_file):
        out = recalc_xlsx(chart_file, "formula_with_chart.xlsx")
        wb = load_workbook(io.BytesIO(out))
        assert wb.sheetnames                 # still a valid workbook
        ws_vals = load_workbook(io.BytesIO(out), data_only=True).active
        assert ws_vals["B1"].value == 6      # SUM(1,2,3)
