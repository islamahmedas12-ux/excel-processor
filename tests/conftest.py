"""
Fixtures للاختبارات
==================
تتضمن ملفات Excel مؤقتة للاستخدام في الاختبارات
"""

import pytest
import os
import tempfile
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Color
import xlwt


@pytest.fixture
def temp_dir():
    """إنشاء دليل مؤقت للاختبارات"""
    with tempfile.TemporaryDirectory() as tmpdir:
        yield tmpdir


@pytest.fixture
def temp_excel_file(temp_dir):
    """إنشاء ملف Excel مؤقت بصيغة .xlsx للاختبارات"""
    file_path = os.path.join(temp_dir, "test_file.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws['A1'] = "Name"
    ws['B1'] = "Age"
    ws['C1'] = "City"

    ws['A2'] = "أحمد"
    ws['B2'] = 25
    ws['C2'] = "الرياض"

    ws['A3'] = "John"
    ws['B3'] = 30
    ws['C3'] = "London"

    cell_a1 = ws['A1']
    cell_a1.font = Font(name='Arial', size=12, bold=True, color='FF0000')
    cell_a1.fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')

    ws['D1'] = 100
    ws['D1'].number_format = '#,##0.00'

    wb.save(file_path)
    return file_path


@pytest.fixture
def temp_old_excel_file(temp_dir):
    """إنشاء ملف Excel مؤقت بصيغة .xls للاختبارات"""
    file_path = os.path.join(temp_dir, "test_old_file.xls")

    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sheet1")

    ws.write(0, 0, "Name")
    ws.write(0, 1, "Age")
    ws.write(0, 2, "City")

    ws.write(1, 0, "أحمد")
    ws.write(1, 1, 25)
    ws.write(1, 2, "الرياض")

    ws.write(2, 0, "John")
    ws.write(2, 1, 30)
    ws.write(2, 2, "London")

    wb.save(file_path)
    return file_path


@pytest.fixture
def temp_excel_with_formulas(temp_dir):
    """إنشاء ملف Excel يحتوي على صيغ"""
    file_path = os.path.join(temp_dir, "test_formulas.xlsx")

    wb = Workbook()
    ws = wb.active

    ws['A1'] = 10
    ws['A2'] = 20
    ws['A3'] = 30
    ws['A4'] = '=SUM(A1:A3)'

    wb.save(file_path)
    return file_path


@pytest.fixture
def temp_multi_sheet_excel(temp_dir):
    """إنشاء ملف Excel متعدد الأوراق"""
    file_path = os.path.join(temp_dir, "test_multi_sheet.xlsx")

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "البيانات"
    ws1['A1'] = "بيانات1"
    ws1['B1'] = "بيانات2"

    ws2 = wb.create_sheet("الإجمالي")
    ws2['A1'] = "الإجمالي"
    ws2['B1'] = 100

    ws3 = wb.create_sheet("Summary")
    ws3['A1'] = "Total"
    ws3['B1'] = 200

    wb.save(file_path)
    return file_path


@pytest.fixture
def temp_excel_empty(temp_dir):
    """إنشاء ملف Excel فارغ"""
    file_path = os.path.join(temp_dir, "test_empty.xlsx")

    wb = Workbook()
    wb.active['A1'] = None

    wb.save(file_path)
    return file_path
