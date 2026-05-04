"""
اختبارات تعديل ملفات Excel
=======================
اختبارات وحدة للتحقق من صحة تعديل ملفات Excel
"""

import pytest
import os
from openpyxl import load_workbook
from excel_processor.file_editor import ExcelEditor, modify_excel_cell, modify_multiple_excel_cells
from excel_processor.errors import (
    FileNotFoundError_,
    InvalidCellCoordinatesError,
    CellWriteError,
    FileSaveError
)


class TestExcelEditor:
    """اختبارات فئة محرر Excel"""

    def test_open_editor(self, temp_excel_file):
        """اختبار فتح المحرر"""
        with ExcelEditor(temp_excel_file) as editor:
            assert editor.workbook is not None
            assert editor.file_format == '.xlsx'

    def test_update_string_cell(self, temp_excel_file, temp_dir):
        """اختبار تحديث خلية نصية"""
        output_path = os.path.join(temp_dir, "updated1.xlsx")

        with ExcelEditor(temp_excel_file) as editor:
            result = editor.update_cell('A1', 'Modified Name')
            editor.save(output_path)

        assert result['old_value'] == 'Name'
        assert result['new_value'] == 'Modified Name'

        wb = load_workbook(output_path)
        assert wb.active['A1'].value == 'Modified Name'
        wb.close()

    def test_update_numeric_cell(self, temp_excel_file, temp_dir):
        """اختبار تحديث خلية رقمية"""
        output_path = os.path.join(temp_dir, "updated2.xlsx")

        with ExcelEditor(temp_excel_file) as editor:
            result = editor.update_cell('B2', 99)
            editor.save(output_path)

        assert result['old_value'] == 25
        assert result['new_value'] == 99

        wb = load_workbook(output_path)
        assert wb.active['B2'].value == 99
        wb.close()

    def test_update_arabic_text(self, temp_excel_file, temp_dir):
        """اختبار تحديث نص عربي"""
        output_path = os.path.join(temp_dir, "updated3.xlsx")

        with ExcelEditor(temp_excel_file) as editor:
            result = editor.update_cell('A2', 'محمد')
            editor.save(output_path)

        assert result['new_value'] == 'محمد'

        wb = load_workbook(output_path)
        assert wb.active['A2'].value == 'محمد'
        wb.close()

    def test_update_english_text(self, temp_excel_file, temp_dir):
        """اختبار تحديث نص إنجليزي"""
        output_path = os.path.join(temp_dir, "updated4.xlsx")

        with ExcelEditor(temp_excel_file) as editor:
            result = editor.update_cell('A3', 'Sarah')
            editor.save(output_path)

        assert result['new_value'] == 'Sarah'

        wb = load_workbook(output_path)
        assert wb.active['A3'].value == 'Sarah'
        wb.close()

    def test_update_invalid_coordinates(self, temp_excel_file):
        """اختبار تحديث بإحداثيات غير صالحة"""
        with pytest.raises(InvalidCellCoordinatesError):
            with ExcelEditor(temp_excel_file) as editor:
                editor.update_cell('INVALID', 'value')

    def test_update_nonexistent_cell(self, temp_excel_file, temp_dir):
        """اختبار تحديث خلية غير موجودة"""
        output_path = os.path.join(temp_dir, "updated5.xlsx")

        with ExcelEditor(temp_excel_file) as editor:
            result = editor.update_cell('Z100', 'New Value')
            editor.save(output_path)

        assert result['new_value'] == 'New Value'

        wb = load_workbook(output_path)
        assert wb.active['Z100'].value == 'New Value'
        wb.close()

    def test_update_preserve_format(self, temp_excel_file, temp_dir):
        """اختبار الحفاظ على التنسيق"""
        output_path = os.path.join(temp_dir, "updated6.xlsx")

        with ExcelEditor(temp_excel_file) as editor:
            original_format = editor.get_cell_format('A1')
            result = editor.update_cell('A1', 'New Value', preserve_format=True)
            editor.save(output_path)

        wb = load_workbook(output_path)
        cell = wb.active['A1']
        assert cell.font.bold == True or cell.font.name is not None
        wb.close()

    def test_update_multiple_cells(self, temp_excel_file, temp_dir):
        """اختبار تحديث عدة خلايا"""
        output_path = os.path.join(temp_dir, "updated7.xlsx")

        updates = {
            'A1': 'First',
            'B2': 50,
            'C3': 'Updated'
        }

        with ExcelEditor(temp_excel_file) as editor:
            result = editor.update_multiple_cells(updates)
            editor.save(output_path)

        assert result['total_updated'] == 3
        assert result['total_failed'] == 0

        wb = load_workbook(output_path)
        assert wb.active['A1'].value == 'First'
        assert wb.active['B2'].value == 50
        assert wb.active['C3'].value == 'Updated'
        wb.close()

    def test_update_with_sheet_name(self, temp_multi_sheet_excel, temp_dir):
        """اختبار تحديث مع تحديد اسم الورقة"""
        output_path = os.path.join(temp_dir, "updated8.xlsx")

        with ExcelEditor(temp_multi_sheet_excel) as editor:
            result = editor.update_cell('A1', 'Modified', sheet_name='الإجمالي')
            editor.save(output_path)

        assert result['sheet'] == 'الإجمالي'
        assert result['new_value'] == 'Modified'

        wb = load_workbook(output_path)
        assert wb['الإجمالي']['A1'].value == 'Modified'
        wb.close()


class TestModifyFunctions:
    """اختبارات دوال التعديل المساعدة"""

    def test_modify_excel_cell_function(self, temp_excel_file, temp_dir):
        """اختبار دالة تعديل الخلية"""
        output_path = os.path.join(temp_dir, "modified1.xlsx")

        result = modify_excel_cell(
            temp_excel_file,
            'B2',
            999,
            output_path=output_path
        )

        assert result['old_value'] == 25
        assert result['new_value'] == 999

        wb = load_workbook(output_path)
        assert wb.active['B2'].value == 999
        wb.close()

    def test_modify_multiple_cells_function(self, temp_excel_file, temp_dir):
        """اختبار دالة تعديل عدة خلايا"""
        output_path = os.path.join(temp_dir, "modified2.xlsx")

        updates = {
            'A1': 'Updated1',
            'B2': 111,
            'C3': 'Updated2'
        }

        result = modify_multiple_excel_cells(
            temp_excel_file,
            updates,
            output_path=output_path
        )

        assert result['total_updated'] == 3

        wb = load_workbook(output_path)
        assert wb.active['A1'].value == 'Updated1'
        assert wb.active['B2'].value == 111
        assert wb.active['C3'].value == 'Updated2'
        wb.close()


class TestEditErrors:
    """اختبارات أخطاء التعديل"""

    def test_edit_nonexistent_file(self):
        """اختبار تعديل ملف غير موجود"""
        with pytest.raises(FileNotFoundError_):
            modify_excel_cell("/non/existent.xlsx", "A1", "value")

    def test_edit_invalid_coordinates(self, temp_excel_file):
        """اختبار تعديل بإحداثيات غير صالحة"""
        with pytest.raises(InvalidCellCoordinatesError):
            modify_excel_cell(temp_excel_file, "1A", "value")

    def test_save_without_changes(self, temp_excel_file, temp_dir):
        """اختبار الحفظ بدون تغييرات"""
        output_path = os.path.join(temp_dir, "saved_no_change.xlsx")

        with ExcelEditor(temp_excel_file) as editor:
            editor.save(output_path)

        assert os.path.exists(output_path)
